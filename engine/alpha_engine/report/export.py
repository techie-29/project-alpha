"""PDF and Excel versions of a finished (v2) report, built only from the report JSON."""
from __future__ import annotations
import io
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

GREEN = colors.HexColor("#0b6b4f")


def fv(v, fmt="number", unit=""):
    if v is None:
        return "-"
    if fmt == "pct":
        return f"{v * 100:.1f}%"
    if fmt == "rating":
        return f"{v:.2f}"
    if fmt == "days":
        return f"{v:.1f} d"
    a = abs(v)
    s = f"{v / 1e6:.2f}M" if a >= 1e6 else f"{v / 1e3:.1f}K" if a >= 1e4 else f"{v:,.0f}" if a >= 100 else f"{v:,.2f}"
    return f"{unit} {s}".strip() if fmt == "money" else s


def _chg(c):
    return "-" if c is None else f"{c * 100:+.1f}%"


def _line(chart, width=170 * mm, height=55 * mm):
    g = chart["grains"].get("month") or chart["grains"][chart["default_grain"]]
    vals = g["series"][0]["values"]
    pts = [(i, v) for i, v in enumerate(vals) if v is not None]
    if len(pts) < 2:
        return None
    d = Drawing(width, height)
    lp = LinePlot()
    lp.x, lp.y, lp.width, lp.height = 30, 18, width - 40, height - 30
    lp.data = [pts]
    lp.lines[0].strokeColor = GREEN
    lp.xValueAxis.valueMin, lp.xValueAxis.valueMax = 0, len(vals) - 1
    step = max(1, len(vals) // 6)
    lp.xValueAxis.valueSteps = list(range(0, len(vals), step))
    lp.xValueAxis.labelTextFormat = lambda i: g["x"][int(i)] if 0 <= int(i) < len(g["x"]) else ""
    lp.xValueAxis.labels.fontSize = lp.yValueAxis.labels.fontSize = 6
    lp.yValueAxis.labelTextFormat = lambda v: fv(v, chart["fmt"] if chart["fmt"] != "money" else "number")
    d.add(lp)
    d.add(String(30, height - 8, chart["title"] + (f" ({chart['unit']})" if chart.get("unit") else ""), fontSize=8, fillColor=colors.black))
    return d


def _bars(chart, width=170 * mm):
    rows = [r for r in chart.get("rows", []) if r.get("value") is not None][:10]
    if len(rows) < 2:
        return None
    h = 24 + 11 * len(rows)
    d = Drawing(width, h)
    bc = HorizontalBarChart()
    bc.x, bc.y, bc.width, bc.height = 90, 6, width - 110, h - 22
    bc.data = [[r["value"] for r in rows][::-1]]
    bc.categoryAxis.categoryNames = [str(r["group"])[:22] for r in rows][::-1]
    bc.categoryAxis.labels.fontSize = bc.valueAxis.labels.fontSize = 6
    bc.valueAxis.labelTextFormat = lambda v: fv(v, chart.get("fmt", "number") if chart.get("fmt") != "money" else "number")
    bc.valueAxis.valueMin = 0
    bc.bars[0].fillColor = GREEN
    d.add(bc)
    d.add(String(90, h - 10, chart["title"] + (f" ({chart['unit']})" if chart.get("unit") else ""), fontSize=8))
    return d


def to_pdf(report: dict, owner: dict | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title="Business report", author="Project Alpha")
    ss = getSampleStyleSheet()
    H1 = ParagraphStyle("H1", parent=ss["Title"], fontSize=19, leading=23, alignment=0, spaceAfter=4)
    H2 = ParagraphStyle("H2", parent=ss["Heading2"], fontSize=13, spaceBefore=10, spaceAfter=5, textColor=GREEN)
    B = ParagraphStyle("B", parent=ss["BodyText"], fontSize=9.2, leading=12.5)
    S = ParagraphStyle("S", parent=B, fontSize=7.8, textColor=colors.HexColor("#555555"))
    T = ParagraphStyle("T", parent=B, fontName="Helvetica-Bold")
    p = lambda t, st=B: Paragraph(escape(str(t)), st)
    b, ov = report["briefing"], report["overview"]
    el = [p("Business report", H1), p(f"Generated {report['generated_at'][:16].replace('T', ' ')} UTC  |  {b.get('coverage') or ''}", S), Spacer(1, 6),
          p(b["headline"], ParagraphStyle("hl", parent=B, fontSize=13, leading=17, fontName="Helvetica-Bold"))]
    el += [p("- " + k) for k in b["key_numbers"]]
    el.append(p("Your files", H2))
    rows = [["File", "What it is", "Rows", "Period"]] + [[f["name"], f["role_text"] + (f" ({f['entity']})" if f.get("entity") else ""), f"{f['rows']:,}" if f.get("rows") else "-",
                                                          f"{f['from']} to {f['to']}" if f.get("from") else "-"] for f in ov["files"]]
    tb = Table(rows, colWidths=[55 * mm, 50 * mm, 22 * mm, 45 * mm])
    tb.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f3ee")), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#bbbbbb"))]))
    el.append(tb)
    el.append(p("Key numbers", H2))
    rows = [["Measure", "Value", "Last month vs previous", "vs a year ago"]]
    for t in ov["tiles"]:
        rows.append([t["label"], fv(t["value"], t["fmt"], t.get("unit") or ""), _chg((t.get("last_period") or {}).get("change")), _chg((t.get("vs_last_year") or {}).get("change"))])
    tb = Table(rows, colWidths=[62 * mm, 40 * mm, 36 * mm, 34 * mm])
    tb.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f3ee")), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#bbbbbb"))]))
    el.append(tb)
    charts = []
    for s in ov["sections"]:
        for c in s["charts"]:
            if c["type"] == "line" and len(charts) < 2:
                charts.append(_line(c))
            elif c["type"] == "bar" and len(charts) < 5 and not c.get("keep_order"):
                charts.append(_bars(c))
    for d in [c for c in charts if c]:
        el += [Spacer(1, 4), d]
    el += [PageBreak(), p("What to do", H2)]
    for key, title in (("fix", "Fix"), ("keep", "Keep"), ("plan", "Plan"), ("watch", "Watch")):
        items = report["suggestions"].get(key, [])
        if items:
            el.append(p(f"{title} ({len(items)})", T))
            for s in items[:8]:
                el += [p("- " + s["title"]), p(s["action"], S)]
    el.append(p("Findings", H2))
    items = [i for i in report["insights"]["items"] if not i.get("muted")][:25]
    for i in items:
        block = [p(f"{i['title']}  [{i['claim']}]", T), p(i["detail"])]
        if i.get("ai_note"):
            block.append(p("AI helper (numbers checked against calculations): " + i["ai_note"], S))
        if i.get("suggestion"):
            block.append(p("Suggestion: " + i["suggestion"], S))
        el.append(KeepTogether(block + [Spacer(1, 5)]))
    if not items:
        el.append(p("Nothing stood out beyond normal variation."))
    answers = {a["questionId"]: a["answer"] for a in (owner or {}).get("answers", [])}
    if report.get("questions"):
        el.append(p("Questions for the owner", H2))
        for q in report["questions"]:
            el += [p(q["text"], S), p("Answer: " + answers.get(q["id"], "(not answered)"))]
    el.append(p("How to read this report", H2))
    el += [p(b["how_to_read"], S)] + [p("- " + l, S) for l in report["limits"]] + [p("- " + c, S) for c in b["caveats"]]

    def foot(c, d):
        c.saveState(); c.setFont("Helvetica", 7.5); c.setFillColor(colors.HexColor("#777777"))
        c.drawString(18 * mm, 9 * mm, "Alpha - shows what the data says and what lines up with what; it does not prove causes."); c.drawRightString(A4[0] - 18 * mm, 9 * mm, f"Page {d.page}"); c.restoreState()
    doc.build(el, onFirstPage=foot, onLaterPages=foot)
    return buf.getvalue()


def _sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B6B4F"); c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([x if not isinstance(x, (list, dict)) else str(x) for x in r])
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def to_xlsx(report: dict, owner: dict | None = None) -> bytes:
    wb = Workbook(); b, ov = report["briefing"], report["overview"]
    ws = wb.active; ws.title = "Briefing"
    _sheet(ws, ["Item", "Text"], [["Headline", b["headline"]], *[["Key number", k] for k in b["key_numbers"]], ["Coverage", b.get("coverage") or ""],
                                  ["Generated", report["generated_at"]], *[["Caveat", c] for c in b["caveats"]], *[["Limit", l] for l in report["limits"]]], [16, 110])
    _sheet(wb.create_sheet("Files"), ["File", "What it is", "Entity", "Rows", "From", "To", "Currency"],
           [[f["name"], f["role_text"], f.get("entity"), f.get("rows"), f.get("from"), f.get("to"), f.get("currency")] for f in ov["files"]], [40, 22, 14, 12, 12, 12, 10])
    _sheet(wb.create_sheet("Key numbers"), ["Measure", "Value", "Unit", "How it is calculated", "Last period", "Change vs previous", "Change vs a year ago"],
           [[t["label"], t["value"], t.get("unit"), t.get("how"), (t.get("last_period") or {}).get("label"), (t.get("last_period") or {}).get("change"), (t.get("vs_last_year") or {}).get("change")] for t in ov["tiles"]],
           [30, 16, 8, 40, 14, 16, 16])
    _sheet(wb.create_sheet("Insights"), ["ID", "Type", "Area", "Claim", "Good or bad", "Title", "Detail", "Suggestion", "Impact", "Unit", "Score", "File", "AI helper note"],
           [[i["id"], i["type"], i["area"], i["claim"], i["polarity"], i["title"], i["detail"], i.get("suggestion"), i.get("impact"), i.get("impact_unit"), i.get("score"), i.get("source"), i.get("ai_note")]
            for i in report["insights"]["items"]], [12, 16, 12, 20, 10, 60, 70, 60, 14, 8, 8, 28, 60])
    _sheet(wb.create_sheet("Suggestions"), ["Bucket", "Finding", "Action", "Claim"],
           [[k.capitalize(), s["title"], s["action"], s["claim"]] for k, v in report["suggestions"].items() for s in v], [10, 60, 70, 20])
    rows = []
    for s in ov["sections"]:
        for c in s["charts"]:
            if c["type"] == "line":
                g = c["grains"].get("month") or c["grains"][c["default_grain"]]
                rows += [[s["title"], c["title"], x, g["series"][0]["name"], v] for x, v in zip(g["x"], g["series"][0]["values"])]
            elif c.get("rows") and c["type"] in ("bar", "histogram"):
                rows += [[s["title"], c["title"], r["group"], c.get("metric") or "", r["value"]] for r in c["rows"]]
            elif c["type"] in ("stacked", "stacked_count"):
                rows += [[s["title"], c["title"], x, se["name"], v] for se in c["series"] for x, v in zip(c["x"], se["values"])]
    _sheet(wb.create_sheet("Chart data"), ["Section", "Chart", "Period / group", "Series", "Value"], rows, [16, 50, 20, 24, 16])
    _sheet(wb.create_sheet("Coverage"), ["File", "Column", "Read as", "Confidence", "What Alpha did"], [[c["file"], c["column"], c["label"], c["confidence"], c["status"]] for c in report["coverage"]], [34, 28, 22, 10, 50])
    ev = []
    for i in report["insights"]["items"][:40]:
        e = i.get("evidence") or {}
        for line in (e.get("rows") or [])[:40]:
            ev.append([i["id"], i["title"][:80], e.get("file"), *line])
    _sheet(wb.create_sheet("Rows behind findings"), ["Finding", "Title", "File"] + [f"Col {k + 1}" for k in range(12)], ev, [12, 50, 28])
    fb = {x["findingId"]: x["verdict"] for x in (owner or {}).get("feedback", [])}
    _sheet(wb.create_sheet("Owner notes"), ["Type", "Item", "Answer or verdict"], [*[["Feedback", k, v] for k, v in fb.items()], *[["Answer", a["questionId"], a["answer"]] for a in (owner or {}).get("answers", [])]], [12, 40, 80])
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()
